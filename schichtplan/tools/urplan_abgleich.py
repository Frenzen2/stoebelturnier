"""Abgleich Urplan ↔ aktueller Excel-Stand für die Übernahme (Version 2).

Statt den aktuellen Excel-Stand als Grundplan zu übernehmen, wird der Urplan zum Grundplan.
Jede Abweichung wird als Ereignis angelegt, wie es die App auch im Betrieb tun würde:

  * Vertretung  (kind 'vacation'): F/S/N einer abwesenden Person (Urlaub/ZA/Krank/FB) macht jemand anderes
  * Übernahme   (kind 'release') : F/S/N einer anwesenden Person macht jemand anderes
  * Tausch      (pair)           : zwei Übernahmen in Gegenrichtung zwischen denselben zwei Personen –
                                   am selben Tag oder innerhalb des Tauschbereichs (von der letzten
                                   regulären F/S/N-Schicht bis zur nächsten, siehe tauschbereich())

TD-Dienste (Platzhalter) sind für Vertretungen gedacht: wer mit TD einen Dienst übernimmt,
hat danach nur diesen Dienst (der TD fällt weg) – genau wie in der App.

Jeder Tag wird zum Schluss nachgerechnet (gleiche Regeln wie effectiveDay() in app.js).
Stimmt das Ergebnis nicht exakt mit dem Excel-Stand überein, bleibt der Tag wie bisher
(Excel-Stand als Grundplan + Markierungen) und kommt auf die Prüfliste.
"""
import datetime
from collections import Counter, defaultdict

import openpyxl

REAL = ('F', 'S', 'N')


def parts(code):
    return [c for c in str(code or '').split('+') if c]


def is_real(c):
    return not str(c).upper().startswith('TD')


def read_urplan(paths, code_of):
    """Urplan-Dateien lesen: Tage als Zeilen, Personen (Nachnamen) als Spalten, ohne Farben."""
    urplan = {}
    for path in paths:
        wb = openpyxl.load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            date_col = next((c for c in range(1, min(ws.max_column, 10) + 1)
                             if any(isinstance(ws.cell(r, c).value, datetime.datetime) for r in range(1, 12))), None)
            if not date_col:
                continue
            header = next((r for r in range(1, 6) if any(isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value.strip()
                                                         for c in range(date_col + 1, date_col + 4))), None)
            if not header:
                continue
            people = []
            for c in range(date_col + 1, ws.max_column + 1):
                v = ws.cell(header, c).value
                if not (isinstance(v, str) and v.strip()):
                    break          # erster leerer Kopf beendet den Personenblock
                people.append((v.strip(), c))
            for r in range(header + 1, ws.max_row + 1):
                d = ws.cell(r, date_col).value
                if not isinstance(d, datetime.datetime):
                    continue
                row = {}
                for name, c in people:
                    code = code_of(ws.cell(r, c).value)
                    row[name] = code or ''   # unbekannter Text (z. B. Schulung) → frei
                urplan[d.strftime('%Y-%m-%d')] = row
    return urplan


def simulate(base, events):
    """Wie effectiveDay() in app.js für einen Tag: base {p: code}, events in Reihenfolge."""
    day = {p: '+'.join(parts(c)) for p, c in base.items()}

    def drop_ph(cur):
        p = parts(cur)
        for i, c in enumerate(p):
            if not is_real(c):
                del p[i]
                break
        return '+'.join(p)

    def remove(cur, code):
        p = parts(cur)
        for c in parts(code):
            if c in p:
                p.remove(c)
        return '+'.join(p)

    for e in events:
        if e['kind'] == 'release':
            day[e['from']] = remove(day.get(e['from'], ''), e['code'])
        day[e['to']] = '+'.join(parts(drop_ph(day.get(e['to'], ''))) + parts(e['code']))
    return day


def same(a, b):
    return Counter(parts(a)) == Counter(parts(b))


def abgleich(out, urplan):
    """Wandelt out (Version 1) in Version 2 um. Gibt (statistik, pruefliste) zurück."""
    plan, marks = out['plan'], out['marks']
    absent = defaultdict(set)
    for a in out['absences']:
        d = datetime.date.fromisoformat(a['from'])
        while d.isoformat() <= a['to']:
            absent[d.isoformat()].add(a['person'])
            d += datetime.timedelta(days=1)
    names = [p['key'] for p in out['people']]
    stats, review = Counter(), []
    events = []          # je Tag gesammelt, Paare werden danach gebildet
    new_plan, new_marks = {}, {}

    for date in sorted(plan):
        cur = plan[date]
        if date not in urplan:
            new_plan[date] = cur
            if date in marks:
                new_marks[date] = marks[date]
            stats['Tage ohne Urplan (Excel-Stand übernommen)'] += 1
            continue
        ur = {n: urplan[date].get(n, '') for n in names}
        A = absent[date]
        # Wer macht welche echten Dienste? (Urplan vs. Excel; Abwesende machen nichts)
        have = {n: Counter(c for c in parts(ur[n]) if is_real(c)) for n in names}
        want = {n: Counter() if n in A else Counter(c for c in parts(cur.get(n, '')) if is_real(c)) for n in names}
        removed = {n: have[n] - want[n] for n in names}
        added = {n: want[n] - have[n] for n in names}
        day_events = []
        for m in names:
            for code, k in list(added[m].items()):
                for _ in range(k):
                    cands = [n for n in names if n != m and removed[n][code] > 0]
                    if not cands:
                        continue
                    hint = (marks.get(date) or {}).get(m)
                    # bevorzugt: laut Excel-Markierung, dann abwesende Person, dann Reihenfolge
                    cands.sort(key=lambda n: (n != hint, n not in A, names.index(n)))
                    n = cands[0]
                    removed[n][code] -= 1
                    added[m][code] -= 1
                    day_events.append({'kind': 'vacation' if n in A else 'release',
                                       'date': date, 'code': code, 'from': n, 'to': m, 'pair': None})
        leftover_add = {n: +added[n] for n in names if +added[n]}
        leftover_rem = {n: +removed[n] for n in names if +removed[n] and n not in A}

        # Platzhalter (TD) wie in der Excel: nur bei Personen ohne übernommenen Dienst anpassen
        base = dict(ur)
        takers = {e['to'] for e in day_events}
        for n in names:
            if n in A or n in takers:
                continue
            want_ph = [c for c in parts(cur.get(n, '')) if not is_real(c)]
            have_ph = [c for c in parts(ur[n]) if not is_real(c)]
            if Counter(want_ph) != Counter(have_ph):
                base[n] = '+'.join([c for c in parts(ur[n]) if is_real(c)] + want_ph)
                stats['TD geändert (direkt im Grundplan)'] += 1

        ok = not leftover_add and not leftover_rem
        if ok:
            for order in (day_events, list(reversed(day_events))):
                res = simulate(base, order)
                if any(not same(res.get(n, ''), cur.get(n, '')) for n in names if n not in A):
                    ok = False
        if not ok:
            # Tag nicht eindeutig → wie bisher: Excel-Stand + Markierungen
            new_plan[date] = cur
            if date in marks:
                new_marks[date] = marks[date]
            details = []
            for n, c in leftover_add.items():
                details.append(f'{n} hat zusätzlich {"+".join(c.elements())} (keine Quelle)')
            for n, c in leftover_rem.items():
                details.append(f'{n} hat {"+".join(c.elements())} nicht mehr (niemand übernimmt)')
            review.append(f'{date}: ' + ('; '.join(details) or 'Ergebnis weicht ab') + ' → Excel-Stand übernommen')
            stats['Tage auf der Prüfliste'] += 1
            continue
        new_plan[date] = base
        events.extend(day_events)
        stats['Tage aus Urplan + Änderungen'] += 1

    # Tausch-Paare: A übernimmt von B und B von A – am selben Tag oder im Tauschbereich
    rel = [e for e in events if e['kind'] == 'release']
    dates_sorted = sorted(urplan)

    def tauschbereich(person, date):
        """Von der letzten regulären F/S/N-Schicht vor dem Tag bis zur nächsten danach (laut Urplan)."""
        i = dates_sorted.index(date)
        lo = hi = date
        for d in reversed(dates_sorted[:i]):
            lo = d
            if any(is_real(c) for c in parts(urplan[d].get(person, ''))):
                break
        for d in dates_sorted[i + 1:]:
            hi = d
            if any(is_real(c) for c in parts(urplan[d].get(person, ''))):
                break
        return lo, hi

    n_pair = 0
    for e in sorted(rel, key=lambda x: x['date']):
        if e['pair']:
            continue
        lo, hi = tauschbereich(e['to'], e['date'])
        lo2, hi2 = tauschbereich(e['from'], e['date'])
        best = None
        for f in rel:
            if f is e or f['pair'] or f['from'] != e['to'] or f['to'] != e['from']:
                continue
            if lo <= f['date'] <= hi or lo2 <= f['date'] <= hi2:
                dist = abs((datetime.date.fromisoformat(f['date']) - datetime.date.fromisoformat(e['date'])).days)
                if best is None or dist < best[0]:
                    best = (dist, f)
        if best:
            n_pair += 1
            e['pair'] = best[1]['pair'] = f'p{n_pair}'
            stats['Tausch am selben Tag' if best[0] == 0 else 'Tausch über mehrere Tage'] += 1
    stats['Vertretung (Urlaub/Krank)'] = sum(e['kind'] == 'vacation' for e in events)
    stats['Übernahme ohne Gegenleistung'] = sum(e['kind'] == 'release' and not e['pair'] for e in events)

    out['version'] = 2
    out['plan'] = new_plan
    out['marks'] = new_marks
    out['events'] = events
    out['review'] = review
    out['urplanRange'] = [min(urplan), max(urplan)] if urplan else None
    return stats, review
