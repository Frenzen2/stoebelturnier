#!/usr/bin/env python3
"""Wandelt den bisherigen Excel-Schichtplan (Tage als Zeilen, Personen als Spalten,
Bedeutung über Zellfarben) in eine Übernahme-Datei (JSON) für die Schichtplan-App um.

Farb-Logik der bisherigen Excel:
  * Jede Person hat eine eigene Farbe (Kopfzeile bzw. häufigste Farbe der Spalte).
  * Zelle in der eigenen Farbe     → Urlaub/ZA (Text = ursprünglicher Dienst)
  * Rote Zelle                     → Krankenstand (Text = ursprünglicher Dienst)
  * Schwarze Zelle, farbige Schrift → Dienst übernommen von der Person dieser Farbe
                                      (Vertretung oder Tausch)

Aufruf:  python3 excel_uebernahme.py Schichtplan.xlsx uebernahme.json
Benötigt: pip install openpyxl
"""
import sys, re, json, colorsys, datetime
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles.colors import COLOR_INDEX

SHIFT_TYPES = [
    {'code': 'F', 'label': 'Frühdienst', 'start': '06:00', 'end': '14:15', 'color': '#bfdbfe', 'placeholder': False},
    {'code': 'S', 'label': 'Spätdienst', 'start': '14:00', 'end': '22:15', 'color': '#fecaca', 'placeholder': False},
    {'code': 'N', 'label': 'Nachtdienst', 'start': '22:00', 'end': '06:15', 'color': '#c7d2fe', 'placeholder': False},
    {'code': 'TD1', 'label': 'Tagdienst 1 (auf Abruf)', 'start': '', 'end': '', 'color': '#bbf7d0', 'placeholder': True},
    {'code': 'TD2', 'label': 'Tagdienst 2 (auf Abruf)', 'start': '', 'end': '', 'color': '#a7f3d0', 'placeholder': True},
    {'code': 'TD3', 'label': 'Tagdienst 3 h Pool', 'start': '', 'end': '', 'color': '#d9f99d', 'placeholder': True},
]
BLACK = {'000000', '060903'}
RED = 'FF0000'


def code_of(text):
    """'6-14,15' → F, '14-22,15' → S, '22-6,15' → N, 'TD1 auf Abruf' → TD1 …; sonst None."""
    t = str(text or '').strip()
    if not t:
        return ''
    m = re.match(r'^(\d{1,2})\s*-\s*(\d{1,2})', t)
    if m:
        return {'6': 'F', '14': 'S', '22': 'N'}.get(m.group(1))
    m = re.match(r'^TD\s*(\d)', t, re.I)
    if m:
        return f'TD{m.group(1)}'
    return None


def main(src, dst):
    wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[0]

    # Theme-Farben auflösen (Excel speichert Farben teils als "Designfarbe + Aufhellung")
    theme = wb.loaded_theme.decode() if wb.loaded_theme else ''
    found = dict((n, a or b) for n, a, b in re.findall(
        r'<a:(dk1|lt1|dk2|lt2|accent\d|hlink|folHlink)>.*?(?:srgbClr val="([0-9A-Fa-f]{6})"|lastClr="([0-9A-Fa-f]{6})")', theme, re.S))
    th = [found.get(k, 'FFFFFF') for k in
          ['lt1', 'dk1', 'lt2', 'dk2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6', 'hlink', 'folHlink']]

    def tint(h, t):
        r, g, b = [int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)]
        H, L, S = colorsys.rgb_to_hls(r, g, b)
        L = L * (1 + t) if t < 0 else L * (1 - t) + t
        r, g, b = colorsys.hls_to_rgb(H, L, S)
        return '%02X%02X%02X' % (round(r * 255), round(g * 255), round(b * 255))

    def col(c):
        if c is None:
            return None
        if c.type == 'rgb':
            return c.rgb[-6:].upper()
        if c.type == 'theme':
            return tint(th[c.theme], c.tint or 0)
        if c.type == 'indexed' and c.indexed < len(COLOR_INDEX):
            return COLOR_INDEX[c.indexed][-6:].upper()
        return None

    def fill(f):
        if getattr(f, 'tagname', '') == 'gradientFill':
            return 'GRAD:' + '/'.join(col(s.color) or '' for s in f.stop)
        if f and f.fill_type == 'solid':
            return col(f.fgColor)
        return None

    def dist(a, b):
        return sum((int(a[i:i + 2], 16) - int(b[i:i + 2], 16)) ** 2 for i in (0, 2, 4)) ** .5

    # Kopfzeile: Datum-Spalte und Personen finden
    header_row = 1
    date_col = next(c for c in range(1, ws.max_column + 1)
                    if any(isinstance(ws.cell(r, c).value, datetime.datetime) for r in range(2, 10)))
    people = []
    for c in range(date_col + 1, ws.max_column + 1):
        name = ws.cell(header_row, c).value
        if name and str(name).strip():
            people.append((str(name).strip(), c))

    # Eigene Farbe(n) je Person: Kopfzeilenfarbe + häufigste Füllfarbe der Spalte
    own = {}
    for name, c in people:
        cnt = Counter()
        for r in range(header_row + 1, ws.max_row + 1):
            f = fill(ws.cell(r, c).fill)
            if f and not f.startswith('GRAD') and f not in BLACK and f != RED and f != 'FFFFFF':
                cnt[f] += 1
        colors = [x for x in [fill(ws.cell(header_row, c).fill)] if x]
        if cnt:
            colors.append(cnt.most_common(1)[0][0])
        own[name] = colors

    def owner_of(color, maxd=70, among=None):
        if not color:
            return None
        best, bd = None, 1e9
        for n, cs in own.items():
            if among is not None and n not in among:
                continue
            for c in cs:
                d = dist(c, color)
                if d < bd:
                    best, bd = n, d
        return best if bd <= maxd else None

    plan = defaultdict(dict)          # date → {person: code}
    absent = defaultdict(dict)        # person → {date: (type, note)}
    marks = defaultdict(dict)         # date → {person: fromPerson}
    stats = Counter()
    unknown_text = Counter()
    unknown_mark = Counter()
    year = None

    for r in range(header_row + 1, ws.max_row + 1):
        d = ws.cell(r, date_col).value
        if not isinstance(d, datetime.datetime):
            continue
        date = d.strftime('%Y-%m-%d')
        year = year or d.year
        # wer ist an diesem Tag abwesend (eigene Farbe / rot)? → bevorzugte Zuordnung der Schriftfarben
        absent_today = set()
        for name, c in people:
            f = fill(ws.cell(r, c).fill)
            if f and (f == RED or (not f.startswith('GRAD') and f not in BLACK and owner_of(f) == name) or f.startswith('GRAD')):
                absent_today.add(name)
        for name, c in people:
            cell = ws.cell(r, c)
            text = cell.value
            code = code_of(text)
            f = fill(cell.fill)
            if code is None:      # Schulung o. Ä. → Fortbildung
                absent[name][date] = ('FB', str(text).strip())
                plan[date][name] = ''
                stats['Fortbildung (Text)'] += 1
                continue
            plan[date][name] = code
            if f and f.startswith('GRAD'):
                if RED in f:
                    absent[name][date] = ('K', '')
                    stats['Krank'] += 1
                else:
                    absent[name][date] = ('U', 'Urlaub/ZA (aus Excel)')
                    stats['Urlaub/ZA'] += 1
            elif f == RED:
                absent[name][date] = ('K', '')
                stats['Krank'] += 1
            elif f in BLACK:
                fc = col(cell.font.color) if cell.font and cell.font.color else None
                frm = owner_of(fc, 120, absent_today - {name}) or owner_of(fc, 100)
                if frm and frm != name:
                    marks[date][name] = frm
                    stats['Übernommen von'] += 1
                else:
                    marks[date][name] = None
                    unknown_mark[fc] += 1
                    stats['Übernommen (Person unklar)'] += 1
            elif f and f != 'FFFFFF' and owner_of(f) == name:
                absent[name][date] = ('U', 'Urlaub/ZA (aus Excel)')
                stats['Urlaub/ZA'] += 1
            elif f and f != 'FFFFFF':
                stats[f'sonstige Farbe {f} ({owner_of(f)})'] += 1
            if code:
                stats['Dienste'] += 1

    # Abwesenheiten zu Zeiträumen zusammenfassen (aufeinanderfolgende Tage, gleiche Art)
    absences = []
    for name, days in absent.items():
        cur = None
        for date in sorted(days):
            typ, note = days[date]
            prev = (datetime.date.fromisoformat(date) - datetime.timedelta(days=1)).isoformat()
            if cur and cur['type'] == typ and cur['note'] == note and cur['to'] == prev:
                cur['to'] = date
            else:
                cur = {'person': name, 'from': date, 'to': date, 'type': typ, 'note': note}
                absences.append(cur)

    out = {
        'format': 'schichtplan-uebernahme', 'version': 1, 'source': src.split('/')[-1],
        'createdAt': datetime.datetime.now().isoformat(timespec='seconds'),
        'people': [{'key': n, 'color': '#' + (own[n][-1] if own[n] else 'E5E7EB').lower()} for n, _ in people],
        'types': SHIFT_TYPES,
        'plan': plan, 'absences': absences, 'marks': marks,
    }
    with open(dst, 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False, indent=0)

    dates = sorted(plan)
    print(f'Zeitraum: {dates[0]} – {dates[-1]} ({len(dates)} Tage), Personen: {", ".join(n for n, _ in people)}')
    for k, v in stats.most_common():
        print(f'  {v:5}  {k}')
    by = Counter(a['type'] for a in absences)
    print('Abwesenheiten (Zeiträume):', dict(by))
    if unknown_mark:
        print('Übernommen, Schriftfarbe keiner Person zuordenbar:', dict(unknown_mark))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
